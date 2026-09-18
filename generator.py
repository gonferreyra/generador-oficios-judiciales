#!/usr/bin/env python3
"""
Generador de oficios - CLI y motor de plantillas
Uso CLI: python generator.py constatacion --data datos.json
Uso Web: importado por server.py
"""

import json
import sys
import zipfile
import io
from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

BASE_DIR = Path(__file__).parent
TEMPLATES_DIR = BASE_DIR / "templates"
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True)

# Texto de exención por honorarios de abogado (Art. 111)
TEXTO_EXENCION_ABOGADO = (
    "Hago saber a Ud. que el presente deberá ser diligenciado SIN PREVIO PAGO DE TASAS Y/O AFOROS, "
    "atento al Art. 111 de la Ley 9459 (Codigo Arancelario para abogados y procuradores de la Provincia de Córdoba), "
    "y deberá comparecer por ante el citado Tribunal a los fines de hacer valer los Derechos que le correspondiere "
    "al Organismo que Ud. dirige, en virtud de tramitarse en autos la subasta del inmueble descripto."
)

# Esquema de campos por tipo de oficio
CAMPOS = {
    "constatacion": {
        "required": [
            "dia", "mes", "anio",
            "autos_caratulados", "juzgado",
            "matricula", "titular",
            "que_se_subasta", "designado",
            "descripcion_inmueble", "nomenclatura_catastral",
            "destinatario_constatacion"
        ],
        "labels": {
            "dia": "Día (ej: 15)",
            "mes": "Mes (ej: junio)",
            "anio": "Año (ej: 2026)",
            "autos_caratulados": "Autos caratulados completos",
            "juzgado": "Juzgado (ej: JUZGADO 1A INST CIV COM 32A NOM)",
            "matricula": "Matrícula (ej: 592.064 (13-04))",
            "titular": "Titular (ej: QUIROGA de VERGAGNI, Mabel Beatriz (100%))",
            "que_se_subasta": "Qué se subasta (ej: derechos y acciones hereditarios sobre el referido inmueble)",
            "designado": "Martillero designado (ej: Carlos R. Ferreyra)",
            "descripcion_inmueble": "Descripción completa del inmueble",
            "nomenclatura_catastral": "Nomenclatura catastral (ej: 13-04-49-01-02-018-019)",
            "destinatario_constatacion": "Destinatario (OFICIAL DE JUSTICIA / JUEZ DE PAZ)"
        }
    },
    "rentas": {
        "required": [
            "dia", "mes", "anio",
            "autos_caratulados", "juzgado",
            "matricula", "titular",
            "descripcion_inmueble", "numero_cuenta"
        ],
        "labels": {
            "dia": "Día (ej: 15)",
            "mes": "Mes (ej: junio)",
            "anio": "Año (ej: 2026)",
            "autos_caratulados": "Autos caratulados completos",
            "juzgado": "Juzgado (ej: Juzgado de 1º Instancia y 32º Nominación)",
            "matricula": "Matrícula (ej: 592.064 (13-04))",
            "titular": "Titular (ej: QUIROGA de VERGAGNI, Mabel Beatriz (100%))",
            "descripcion_inmueble": "Descripción completa del inmueble",
            "numero_cuenta": "Número de cuenta (ej: 1304-02708973)"
        }
    },
    "municipalidad": {
        "required": [
            "dia", "mes", "anio",
            "autos_caratulados", "juzgado",
            "matricula", "titular",
            "descripcion_inmueble", "nomenclatura_catastral",
            "ciudad"
        ],
        "labels": {
            "dia": "Día (ej: 15)",
            "mes": "Mes (ej: junio)",
            "anio": "Año (ej: 2026)",
            "autos_caratulados": "Autos caratulados completos",
            "juzgado": "Juzgado (ej: Juzgado de 1º Instancia y 32º Nominación)",
            "matricula": "Matrícula (ej: 592.064 (13-04))",
            "titular": "Titular (ej: QUIROGA de VERGAGNI, Mabel Beatriz (100%))",
            "descripcion_inmueble": "Descripción completa del inmueble",
            "nomenclatura_catastral": "Nomenclatura catastral (ej: 13-04-49-01-02-018-019)",
            "ciudad": "Ciudad/Intendencia (ej: Unquillo)"
        }
    }
}

# Campos globales del formulario (exención de pago - aplica a los 3 oficios)
CAMPOS_GLOBALES = {
    "required": [
        "exento_pago", "tipo_exencion", "texto_exencion_otros"
    ],
    "labels": {
        "exento_pago": "Exento de pago",
        "tipo_exencion": "Tipo de exención",
        "texto_exencion_otros": "Texto de exención personalizada"
    }
}


def obtener_texto_exencion(datos: dict) -> str:
    """Devuelve el texto de exención según los datos del formulario"""
    if not datos.get("exento_pago"):
        return ""
    
    tipo = datos.get("tipo_exencion", "")
    if tipo == "honorario_abogado":
        return TEXTO_EXENCION_ABOGADO
    elif tipo == "otros":
        return datos.get("texto_exencion_otros", "").strip()
    return ""


def cargar_plantilla(tipo: str, firma: str = "tribunal") -> Document:
    """Carga la plantilla .docx base"""
    path = TEMPLATES_DIR / tipo / f"{firma}.docx"
    if not path.exists():
        raise FileNotFoundError(f"Plantilla no encontrada: {path}")
    return Document(path)


def reemplazar_placeholders(doc: Document, datos: dict) -> Document:
    """Reemplaza {{campo}} por valor en párrafos y tablas"""
    for p in doc.paragraphs:
        for run in p.runs:
            for k, v in datos.items():
                if f"{{{{{k}}}}}" in run.text:
                    run.text = run.text.replace(f"{{{{{k}}}}}", str(v))

    # También en tablas
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        for k, v in datos.items():
                            if f"{{{{{k}}}}}" in run.text:
                                run.text = run.text.replace(f"{{{{{k}}}}}", str(v))
    
    # Eliminar párrafos vacíos para evitar espacios extra
    _eliminar_parrafos_vacios(doc)
    return doc


def _eliminar_parrafos_vacios(doc: Document):
    """Elimina párrafos que solo contienen espacios en blanco"""
    for p in list(doc.paragraphs):
        if not p.text.strip():
            p._element.getparent().remove(p._element)


def generar(tipo: str, datos: dict, firma: str = "tribunal", output_name: str = None) -> Path:
    """Genera el oficio y devuelve la ruta del archivo"""
    if tipo not in CAMPOS:
        raise ValueError(f"Tipo de oficio no soportado: {tipo}")

    # Validar campos requeridos
    faltantes = [c for c in CAMPOS[tipo]["required"] if c not in datos]
    if faltantes:
        raise ValueError(f"Faltan campos requeridos: {', '.join(faltantes)}")

    doc = cargar_plantilla(tipo, firma)
    doc = reemplazar_placeholders(doc, datos)

    if not output_name:
        output_name = f"{tipo}_{firma}_{datos.get('matricula', 'sin_matricula').replace('.', '_').replace('/', '_')}.docx"

    output_path = OUTPUT_DIR / output_name
    doc.save(output_path)
    return output_path


def generar_todos(datos: dict, firma: str = "tribunal") -> bytes:
    """Genera los 3 oficios y los devuelve como ZIP en memoria"""
    tipos = ["constatacion", "rentas", "municipalidad"]
    zip_buffer = io.BytesIO()
    
    # Obtener texto de exención (aplica a los 3 oficios)
    texto_exencion = obtener_texto_exencion(datos)
    
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zipf:
        for tipo in tipos:
            # Filtrar solo los campos requeridos para este tipo
            campos_requeridos = CAMPOS[tipo]["required"]
            datos_tipo = {k: v for k, v in datos.items() if k in campos_requeridos}
            
            # Agregar texto de exención (siempre, vacío si no aplica - se eliminará el párrafo vacío)
            datos_tipo["texto_exencion"] = texto_exencion
            
            # Validar campos requeridos
            faltantes = [c for c in campos_requeridos if c not in datos_tipo]
            if faltantes:
                raise ValueError(f"Para {tipo} faltan campos: {', '.join(faltantes)}")
            
            doc = cargar_plantilla(tipo, firma)
            doc = reemplazar_placeholders(doc, datos_tipo)
            
            output_name = f"{tipo}_{firma}_{datos.get('matricula', 'sin_matricula').replace('.', '_').replace('/', '_')}.docx"
            
            doc_buffer = io.BytesIO()
            doc.save(doc_buffer)
            doc_buffer.seek(0)
            
            zipf.writestr(output_name, doc_buffer.read())
    
    zip_buffer.seek(0)
    return zip_buffer.read()


def formulario_interactivo(tipo: str) -> dict:
    """Pide los campos por consola"""
    print(f"\n=== Generando {tipo.upper()} ===\n")
    datos = {}
    for campo in CAMPOS[tipo]["required"]:
        label = CAMPOS[tipo]["labels"].get(campo, campo)
        valor = input(f"{label}: ").strip()
        while not valor:
            valor = input(f"  (requerido) {label}: ").strip()
        datos[campo] = valor
    return datos


def main():
    if len(sys.argv) < 2:
        print("Uso: python generator.py <tipo> [--data archivo.json] [--firma tribunal|martillero] [--output nombre.docx]")
        print("Tipos disponibles:", ", ".join(CAMPOS.keys()))
        sys.exit(1)

    tipo = sys.argv[1]
    firma = "tribunal"
    data_file = None
    output_name = None

    i = 2
    while i < len(sys.argv):
        if sys.argv[i] == "--firma" and i + 1 < len(sys.argv):
            firma = sys.argv[i + 1]
            i += 2
        elif sys.argv[i] == "--data" and i + 1 < len(sys.argv):
            data_file = sys.argv[i + 1]
            i += 2
        elif sys.argv[i] == "--output" and i + 1 < len(sys.argv):
            output_name = sys.argv[i + 1]
            i += 2
        else:
            i += 1

    if data_file:
        with open(data_file) as f:
            datos = json.load(f)
    else:
        datos = formulario_interactivo(tipo)

    try:
        output_path = generar(tipo, datos, firma, output_name)
        print(f"\n✅ Oficio generado: {output_path}")
    except Exception as e:
        print(f"\n❌ Error: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()