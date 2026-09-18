#!/usr/bin/env python3
"""
Recrea las plantillas de rentas y municipalidad con estructura EXACTA
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pathlib import Path

def create_base_doc():
    doc = Document()
    
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)
    
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial Narrow'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)
    
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = Pt(14)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style.paragraph_format.first_line_indent = Cm(1.25)
    
    style.element.rPr.rFonts.set(qn('w:ascii'), 'Arial Narrow')
    style.element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial Narrow')
    style.element.rPr.rFonts.set(qn('w:cs'), 'Arial Narrow')
    
    return doc

def add_para(doc, text, bold=False, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, font_name='Arial Narrow', font_size=Pt(12), indent=True):
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(14)
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = font_size
    run.bold = bold
    return p

def add_mixed_para(doc, parts, alignment=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True):
    p = doc.add_paragraph()
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(14)
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    for part in parts:
        text = part[0]
        bold = part[1] if len(part) > 1 else False
        font_name = part[2] if len(part) > 2 else 'Arial Narrow'
        font_size = part[3] if len(part) > 3 else Pt(12)
        run = p.add_run(text)
        run.font.name = font_name
        run.font.size = font_size
        run.bold = bold
    return p

# ============ MUNICIPALIDAD ============
# Estructura según PDF:
# P1: Fecha (DERECHA)
# P2: Destinatario (IZQUIERDA) - "Sr. Intendente Municipal de la Ciudad De {{ciudad}}"
# P3: Cuerpo principal (CON sangría) - solicita INFORMAR sobre SUBSISTENCIA DE DEUDA FISCAL
#     Incluye descripción del inmueble dentro del cuerpo
# P4: Nomenclatura Catastral (CON sangría, 14-16pt + negrita)
# P5: Texto legal Art. 111 / exención (CON sangría, condicional)
# P6: Cierre (CON sangría)
# P7: DIOS GUARDE A UD (CON sangría, JUSTIFICADO)
def create_municipalidad_template():
    doc = create_base_doc()
    
    # P1: Fecha (DERECHA, SIN sangría)
    add_para(doc, "Córdoba, {{dia}} de {{mes}} de {{anio}}.", bold=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
    
    # P2: Destinatario (IZQUIERDA, SIN sangría)
    add_para(doc, "Sr. Intendente Municipal de la\nCiudad De {{ciudad}}\nS/D", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT, indent=False)
    
    # P3: Cuerpo principal (CON sangría) - INFORMAR sobre SUBSISTENCIA DE DEUDA FISCAL
    add_mixed_para(doc, [
        ('En los autos caratulados "', False),
        ('{{autos_caratulados}}', True),
        ('", que se tramitan por ante este {{juzgado}} de esta Ciudad, secretaría a cargo de la autorizante, se ha resuelto librar a Ud. el presente a fin de que tan luego de recibido y previas las formalidades de Ley, proceda a INFORMAR sobre SUBSISTENCIA DE DEUDA FISCAL en ese Organismo a su cargo, sobre el inmueble que se ejecuta en autos, el que se halla inscripto en el Registro General de la Provincia a la Matrícula {{matricula}} a nombre de {{titular}}; que a continuación se describe: {{descripcion_inmueble}}', False),
    ])
    
    # P4: Nomenclatura Catastral (CON sangría, 14pt + negrita)
    add_mixed_para(doc, [
        ('NOM. CATASTRAL: ', True, 'Arial Narrow', Pt(14)),
        ('{{nomenclatura_catastral}}.-', False, 'Arial Narrow', Pt(14)),
    ])
    
    # P5: Texto de exención / Art. 111 (CON sangría, condicional)
    add_para(doc, '{{texto_exencion}}')
    
    # P6: Cierre (CON sangría)
    add_para(doc, "A la espera del pronto diligenciamiento del presente, saludo a UD. atte.-")
    
    # P7: DIOS GUARDE A UD (CON sangría, JUSTIFICADO)
    add_para(doc, "DIOS GUARDE A UD.-", alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    output_path = Path("templates/municipalidad/tribunal.docx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"✅ Municipalidad: {output_path} ({len(doc.paragraphs)} párrafos)")

# ============ RENTAS ============
# Estructura según PDF:
# P1: Fecha (DERECHA)
# P2: Destinatario (IZQUIERDA)
# P3: Cuerpo principal (CON sangría) - INFORMAR BASE IMPONIBLE + SUBSISTENCIA DE DEUDA FISCAL
#     Incluye descripción del inmueble dentro del cuerpo
# P4: Número de Cuenta (CON sangría, 14pt + negrita, párrafo separado)
# P5: Texto de exención (CON sangría, condicional)
# P6: Cierre (CON sangría)
# P7: DIOS GUARDE A UD (CON sangría, JUSTIFICADO)
def create_rentas_template():
    doc = create_base_doc()
    
    # P1: Fecha (DERECHA, SIN sangría)
    add_para(doc, "Córdoba, {{dia}} de {{mes}} de {{anio}}.", bold=True, alignment=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
    
    # P2: Destinatario (IZQUIERDA, SIN sangría)
    add_para(doc, "Sr. Director de la\nDIRECCION GENERAL DE RENTAS\nS/D", bold=True, alignment=WD_ALIGN_PARAGRAPH.LEFT, indent=False)
    
    # P3: Cuerpo principal (CON sangría) - INFORMAR BASE IMPONIBLE + SUBSISTENCIA DE DEUDA FISCAL
    add_mixed_para(doc, [
        ('En los autos caratulados "', False),
        ('{{autos_caratulados}}', True),
        ('", que se tramitan por ante este {{juzgado}} de esta Ciudad, secretaría a cargo de la autorizante, se ha resuelto librar a Ud. el presente a fin de que tan luego de recibido y previas las formalidades de Ley, proceda a INFORMAR sobre BASE IMPONIBLE, correspondiente al año en curso, como así también sobre SUBSISTENCIA DE DEUDA FISCAL, en ese Organismo a su cargo, sobre el inmueble que se ejecuta en autos, el que se halla inscripto en el Registro General de la Provincia a la Matrícula {{matricula}} a nombre de {{titular}}; que a continuación se describe: {{descripcion_inmueble}}', False),
    ])
    
    # P4: Número de Cuenta (CON sangría, 14pt + negrita, párrafo separado)
    add_mixed_para(doc, [
        ('Número de Cuenta: ', True, 'Arial Narrow', Pt(14)),
        ('{{numero_cuenta}}.-', False, 'Arial Narrow', Pt(14)),
    ])
    
    # P5: Texto de exención (CON sangría, condicional)
    add_para(doc, '{{texto_exencion}}')
    
    # P6: Cierre (CON sangría)
    add_para(doc, "A la espera del pronto diligenciamiento del presente, saludo a UD. atte.-")
    
    # P7: DIOS GUARDE A UD (CON sangría, JUSTIFICADO)
    add_para(doc, "DIOS GUARDE A UD.", alignment=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    output_path = Path("templates/rentas/tribunal.docx")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"✅ Rentas: {output_path} ({len(doc.paragraphs)} párrafos)")

if __name__ == "__main__":
    create_municipalidad_template()
    create_rentas_template()